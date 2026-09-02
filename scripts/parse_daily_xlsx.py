#!/usr/bin/env python3
"""
parse_daily_xlsx.py — Parser determinístico dos relatórios diários FrontOffice
(layout novo, Excel 2007+/.xlsx) para Victory Business (VB) e Victory Suítes (VS).

USO:
    python3 parse_daily_xlsx.py <caminho_VB.xlsx> <caminho_VS.xlsx> [--out saida.json]

Mesma saída que parse_daily.py (dias_efetivo, dias_previsao, subtotal_efetivo,
subtotal_previsao, totalizacao, validacao), mas para o layout novo, que não tem
os marcadores de texto "Efetivo"/"Previsão": os dias são classificados
comparando a data da linha com a data de geração do relatório, embutida na
última célula do rodapé (coluna 24 da última linha não vazia) — data <= geração
= Efetivo (já em andamento/fechado); data > geração = Previsão (futuro/OTB).

NÃO usa LLM, não abre sub-agentes, não imprime dados brutos de célula.
"""
import sys
import json
import argparse
import openpyxl
from datetime import datetime

COL_DATA = 1
COL_PCT_OCC = 4
COL_OCUPADAS = 5
COL_REC_EF = 21
COL_REC_PREV = 22
COL_TOT_REC = 23
COL_ADR = 24


def to_float(v):
    if v == "" or v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(".", "").replace(",", ".") if "," in str(v) else str(v).strip()
    try:
        return float(s)
    except ValueError:
        return 0.0


def find_generation_datetime(sh):
    for r in range(sh.max_row, 0, -1):
        v = sh.cell(row=r, column=24).value
        if isinstance(v, datetime):
            return v
    return None


def _row_to_day(sh, r):
    return {
        "data": sh.cell(row=r, column=COL_DATA).value.strftime("%Y-%m-%d"),
        "dow": sh.cell(row=r, column=2).value,
        "occ_pct": round(to_float(sh.cell(row=r, column=COL_PCT_OCC).value), 2),
        "ocupadas": int(to_float(sh.cell(row=r, column=COL_OCUPADAS).value)),
        "rec_efetivo": round(to_float(sh.cell(row=r, column=COL_REC_EF).value), 2),
        "rec_previsto": round(to_float(sh.cell(row=r, column=COL_REC_PREV).value), 2),
        "rec_total": round(to_float(sh.cell(row=r, column=COL_TOT_REC).value), 2),
        "adr": round(to_float(sh.cell(row=r, column=COL_ADR).value), 2),
    }


def _sum_bucket(days):
    if not days:
        return None
    return {
        "occ_pct": round(sum(d["occ_pct"] for d in days) / len(days), 2) if False else None,
        "ocupadas": sum(d["ocupadas"] for d in days),
        "rec_efetivo": round(sum(d["rec_efetivo"] for d in days), 2),
        "rec_previsto": round(sum(d["rec_previsto"] for d in days), 2),
        "rec_total": round(sum(d["rec_total"] for d in days), 2),
    }


def parse_report(path, room_count):
    wb = openpyxl.load_workbook(path, data_only=True)
    sh = wb.active

    gen_dt = find_generation_datetime(sh)
    if gen_dt is None:
        raise ValueError(f"{path}: não encontrei a data de geração no rodapé (coluna 24)")
    gen_date = gen_dt.date()

    dias_efetivo = []
    dias_previsao = []
    totalizacao = None

    for r in range(1, sh.max_row + 1):
        col0 = sh.cell(row=r, column=COL_DATA).value
        if col0 == "TOTALIZAÇÃO":
            totalizacao = {
                "occ_pct": round(to_float(sh.cell(row=r, column=COL_PCT_OCC).value), 2),
                "ocupadas": int(to_float(sh.cell(row=r, column=COL_OCUPADAS).value)),
                "rec_efetivo": round(to_float(sh.cell(row=r, column=COL_REC_EF).value), 2),
                "rec_previsto": round(to_float(sh.cell(row=r, column=COL_REC_PREV).value), 2),
                "rec_total": round(to_float(sh.cell(row=r, column=COL_TOT_REC).value), 2),
                "adr": round(to_float(sh.cell(row=r, column=COL_ADR).value), 2),
            }
            continue
        if isinstance(col0, datetime):
            day = _row_to_day(sh, r)
            if col0.date() <= gen_date:
                dias_efetivo.append(day)
            else:
                dias_previsao.append(day)

    se = _sum_bucket(dias_efetivo)
    sp = _sum_bucket(dias_previsao)
    if se is not None:
        se["occ_pct"] = round(sum(d["ocupadas"] for d in dias_efetivo) / (room_count * max(len(dias_efetivo), 1)) * 100, 2)
        se["adr"] = round(se["rec_total"] / se["ocupadas"], 2) if se["ocupadas"] else 0.0
    if sp is not None:
        sp["occ_pct"] = round(sum(d["ocupadas"] for d in dias_previsao) / (room_count * max(len(dias_previsao), 1)) * 100, 2)
        sp["adr"] = round(sp["rec_total"] / sp["ocupadas"], 2) if sp["ocupadas"] else 0.0

    validacao = {"ok": None, "detalhe": None}
    if se and sp and totalizacao:
        soma_receita = round(se["rec_total"] + sp["rec_total"], 2)
        diff = round(soma_receita - totalizacao["rec_total"], 2)
        validacao["ok"] = abs(diff) < 1.0
        validacao["detalhe"] = (
            f"Subtotal Ef. ({se['rec_total']:.2f}) + Subtotal Prev. ({sp['rec_total']:.2f}) "
            f"= {soma_receita:.2f} vs TOTALIZAÇÃO ({totalizacao['rec_total']:.2f}) | diff={diff:.2f}"
        )

    return {
        "room_count": room_count,
        "geracao": gen_dt.isoformat(),
        "dias_efetivo": dias_efetivo,
        "dias_previsao": dias_previsao,
        "subtotal_efetivo": se,
        "subtotal_previsao": sp,
        "totalizacao": totalizacao,
        "validacao": validacao,
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("vb_path")
    ap.add_argument("vs_path")
    ap.add_argument("--out", default="daily_data_xlsx.json")
    args = ap.parse_args()

    result = {
        "gerado_em": datetime.now().isoformat(timespec="seconds"),
        "VB": parse_report(args.vb_path, room_count=152),
        "VS": parse_report(args.vs_path, room_count=122),
    }

    with open(args.out, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    for hotel in ("VB", "VS"):
        v = result[hotel]["validacao"]
        print(f"{hotel}: {len(result[hotel]['dias_efetivo'])} dias efetivo, "
              f"{len(result[hotel]['dias_previsao'])} dias previsão | "
              f"geração={result[hotel]['geracao']} | validação={'OK' if v['ok'] else 'FALHOU'}")
        if not v["ok"]:
            print(f"  -> {v['detalhe']}")

    print(f"\nJSON salvo em: {args.out}")


if __name__ == "__main__":
    main()
